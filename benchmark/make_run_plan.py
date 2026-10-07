"""Write benchmark/test_plan.xlsx: the ordered, randomised run list for the insert-mode study.

    python benchmark/make_run_plan.py

The order inside each phase is shuffled with a fixed seed (SEED), so the list is reproducible
and the same seed is reported in the paper. Phases follow the dependency chain of the plan:
T1 picks the clearance (C*), T4 picks the adhesion recipe (R*), T3 then uses both.
"""

from __future__ import annotations

import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 20261006
OUT = Path(__file__).resolve().parent / "test_plan.xlsx"

MODEL = {
    "M1": ("M1_tensile_D638-I_t7", 82.5, 412),
    "M2": ("M2_flex_80x10x5", 40.0, 200),
    "G1": ("G1_square_20", 20.0, 100), "G2": ("G2_cylinder_d20", 20.0, 100),
    "G3a": ("G3a_cone_10deg", 20.0, 100), "G3b": ("G3b_cone_25deg", 20.0, 100),
    "G4": ("G4_tube_d24_d16", 20.0, 100), "G5": ("G5_L_30x20", 20.0, 100),
    "G6": ("G6_thin_5x30", 3.0, 15), "G7": ("G7_overhang", 20.0, 100),
}
CSTAR = "C*"   # best clearance from T1 phase A (filled in on the Parametreler sheet)
RSTAR = "R*"   # best adhesion recipe from T4

# T4: 2^(4-1) fractional factorial, E = ABC (resolution IV). D (reheat pass) waits for the feature.
T4_LEVELS = {"A": (0, 20), "B": (100, 40), "C": ("on", "off"), "E": (0, 10)}
T4_RUNS = []
for i, (a, b, c) in enumerate([(a, b, c) for a in (-1, 1) for b in (-1, 1) for c in (-1, 1)], start=1):
    e = a * b * c
    lv = lambda f, s: T4_LEVELS[f][0 if s < 0 else 1]
    T4_RUNS.append(dict(run=f"r{i:02d}", A=lv("A", a), B=lv("B", b), C=lv("C", c), E=lv("E", e)))


def ins_cmd(model, sid, layer, clearance="0.15", wall=None, extra=""):
    w = f" --wall-height {wall}" if wall is not None else ""
    return (f"python benchmark/run_insert.py {MODEL[model][0]}.gcode -o {sid}.gcode "
            f"--part-layer {layer} --clearance {clearance}{w}{extra}").strip()


BOTTOM = "Alt parça: Bambu Studio'da katman {n}'e Pause ekle, durunca baskıyı iptal et, soğuyunca sök."


def rows():
    out = []

    def add(phase, test, model, group, material, variable, sid, clearance="", wall="", pset="", cmd="", manual="", seam=None, layer=None):
        m = MODEL.get(model)
        out.append(dict(phase=phase, test=test, model=(m[0] + ".stl") if m else model, group=group, material=material,
                        variable=variable, sid=sid, seam=seam if seam is not None else (m[1] if m else ""),
                        layer=layer if layer is not None else (m[2] if m else ""), clearance=clearance, wall=wall,
                        pset=pset, cmd=cmd, manual=manual))

    # Phase 0: pilot
    for k in range(1, 4):
        add("F0 Pilot", "Pilot", "M1", "MONO", "PLA", "CV tahmini", f"P0-M1-MONO-PLA-{k:02d}", pset="S-M",
            manual="Tek seferde bas, 10 mm brim, düşük hızlanma.")
    for k in range(1, 7):
        sid = f"P0-G1-c015-{k:02d}"
        add("F0 Pilot", "Pilot", "G1", "INS-D", "PLA", "ölçüm tekrarlanabilirliği (her parça 3 kez ölçülür)", sid,
            "0.15", "", "S-G", ins_cmd("G1", sid, 100), BOTTOM.format(n=101))

    # Phase 1: T1 A + T2
    for g in ["G1", "G2", "G3a", "G4", "G5", "G6"]:
        for c in ["0.05", "0.10", "0.15", "0.20", "0.30"]:
            for k in range(1, 4):
                sid = f"T1A-{g}-c{c.replace('0.', '0')}-{k:02d}"
                add("F1 T1A+T2", "T1A", g, "INS-D", "PLA", f"boşluk {c} mm", sid, c, "", "S-G",
                    ins_cmd(g, sid, MODEL[g][2], c), BOTTOM.format(n=MODEL[g][2] + 1))
    for g in ["G1", "G2", "G3a", "G4"]:
        add("F1 T1A+T2", "T2", g, "Alt parça", "PLA", "tahribatsız tutma ölçümleri için", f"T2-{g}-part", "", "", "S-G",
            "", BOTTOM.format(n=101) + " Tüm T2 duvarlarında aynı parça kullanılır.")
        for c in ["0.05", "0.10", "0.15", "0.20", "0.30"]:
            sid = f"T2-{g}-wall-c{c.replace('0.', '0')}"
            add("F1 T1A+T2", "T2", g, "Yalnız duvar", "PLA", f"boşluk {c} mm; 5 kuvvet ölçümü", sid, c, "", "S-G",
                ins_cmd(g, sid, 100, c), "Duvar bitip yazıcı durunca baskıyı iptal et. Yanal ve dikey kuvveti 5'er kez ölç.")
    for g in ["G3a", "G3b"]:
        for glue in ["yok", "4 nokta", "2 nokta"]:
            for k in range(1, 4):
                sid = f"T2-{g}-glue{glue[0]}-{k:02d}"
                add("F1 T1A+T2", "T2", g, "INS-D", "PLA", f"sıcak tutkal: {glue}", sid, "0.15", "", "S-G",
                    ins_cmd(g, sid, 100, "0.15"),
                    BOTTOM.format(n=101) + (f" Parçayı oturttuktan sonra duvarın üstüne {glue} sıcak tutkal." if glue != "yok" else ""))

    # Phase 2: T1 B + ablation
    for g in ["G1", "G2", "G5"]:
        for w in ["3", "10", "19"]:
            for k in range(1, 4):
                sid = f"T1B-{g}-w{w}-{k:02d}"
                add("F2 T1B+MAN", "T1B", g, "INS-D", "PLA", f"duvar {w} mm", sid, CSTAR, w, "S-G",
                    ins_cmd(g, sid, 100, "<C*>", w), BOTTOM.format(n=101))
        for op in ["O1", "O2"]:
            for k in range(1, 4):
                sid = f"MAN-{g}-{op}-{k:02d}"
                add("F2 T1B+MAN", "T1-ablasyon", g, "MAN", "PLA", f"duvarsız, operatör {op}", sid, "", "", "S-G",
                    "AÇIK İŞ: duvarsız sürüm yok. run_insert/insert.py'ye duvarı atlayan bir seçenek eklenince: "
                    + ins_cmd(g, sid, 100, "<C*>") + " --no-wall",
                    BOTTOM.format(n=101) + " Parçayı gözle ve kumpasla hizala, bantla sabitle.")

    # Phase 3: T4
    for r in T4_RUNS:
        for k in range(1, 4):
            sid = f"T4-{r['run']}-PLA-{k:02d}"
            fan = " --fan-on" if r["C"] == "on" else " --fan-off"
            extra = f" --temp-boost {r['A']} --speed {r['B']}{fan}"
            add("F3 T4", "T4", "M2", "INS-D", "PLA", f"A +{r['A']} °C · B %{r['B']} · C fan {r['C']} · E {r['E']} dk",
                sid, CSTAR, "", f"S-M / {r['run']}", ins_cmd("M2", sid, 200, "<C*>", extra=extra),
                BOTTOM.format(n=201) + f" Parçayı oturttuktan sonra devam etmeden {r['E']} dk bekle.")

    # Phase 4: T3 (+ T4 validation)
    for mat in ["PLA", "PETG"]:
        for grp in ["MONO", "RES", "INS-D", "INS-B", "GLU-CA", "GLU-EP"]:
            for k in range(1, 8):
                sid = f"T3-{grp}-{mat}-{k:02d}"
                cmd, manual = "", ""
                if grp == "MONO":
                    manual = "Tek seferde bas, 10 mm brim, düşük hızlanma."
                elif grp == "RES":
                    cmd = "layer-rescue M1_tensile_D638-I_t7.gcode --last-layer 412 --z-mode <retained|manual>"
                    manual = ("Katman 413'e Pause ekle, durunca iptal et, parça tablada kalsın, 15 dk bekle, resume G-code'unu bas. "
                              "Z modunu README'deki resume adımlarına göre seç ve Not'a yaz.")
                elif grp == "INS-D":
                    cmd = ins_cmd("M1", sid, 412, "<C*>") + " <R* bayrakları>"
                    manual = BOTTOM.format(n=413)
                elif grp == "INS-B":
                    cmd = ins_cmd("M1", sid, 412, "<C*>") + " <R* bayrakları>"
                    manual = "Tam bas, Z ≈ 83 mm'den kes, 400 kum zımparayla 82,5 mm'ye indir (kumpasla), sonra insert."
                else:
                    glue = "CA" if grp == "GLU-CA" else "epoksi"
                    manual = f"Bambu Studio Cut ile Z = 82,5 mm'den ikiye böl, iki yarıyı ayrı bas, {glue} ile yapıştır, 24 sa beklet."
                add("F4 T3", "T3", "M1", grp, mat, grp, sid, CSTAR if grp.startswith("INS") else "", "",
                    "S-M" + (" / R*" if grp.startswith("INS") else ""), cmd, manual)
    for rec in ["R*", "varsayılan"]:
        for k in range(1, 6):
            sid = f"T4V-{'best' if rec == 'R*' else 'dflt'}-PLA-{k:02d}"
            extra = " <R* bayrakları>" if rec == "R*" else ""
            add("F4 T3", "T4-doğrulama", "M1", "INS-D", "PLA", f"reçete {rec}", sid, CSTAR, "", "S-M",
                ins_cmd("M1", sid, 412, "<C*>") + extra, BOTTOM.format(n=413))

    # Phase 5: T5, T6/T8, T7
    for e in [-3, -2, -1, 0, 1, 2, 3]:
        for model, n, last in [("M1", 3, 412), ("G1", 2, 100)]:
            for k in range(1, n + 1):
                sid = f"T5-{model}-e{e:+d}-{k:02d}"
                add("F5 T5–T8", "T5", model, "INS-D", "PLA", f"girilen katman hatası {e:+d}", sid, CSTAR, "",
                    "S-M" if model == "M1" else "S-G", ins_cmd(model, sid, last + e, "<C*>"),
                    BOTTOM.format(n=last + 1) + f" Fiziksel parça {last}. katmanda; G-code {last + e}. katmana göre üretilir.")
    for r, name, seam in [("R1", "3DBenchy", "21 mm (Test 3)"), ("R2", "SpeedDrone roket", "katman 260 (Test 4)"),
                          ("R3", "Talon burun (LW-PLA)", "belirlenecek"), ("R4", "Mystic Dragon", "belirlenecek")]:
        for k in range(1, 4):
            sid = f"T6-{r}-INS-{k:02d}"
            add("F5 T5–T8", "T6/T8", f"{r} {name}", "INS-D", "PLA" if r != "R3" else "LW-PLA", "gerçek model", sid,
                CSTAR, "", "S-R", f"python benchmark/run_insert.py {r}.gcode -o {sid}.gcode --part-height/--part-layer <dikiş> --clearance <C*>",
                "Alt parçayı dikişte durdur, sök, insert. 3B tara.", seam=seam, layer="")
        add("F5 T5–T8", "T6/T8", f"{r} {name}", "MONO", "PLA" if r != "R3" else "LW-PLA", "referans tarama", f"T6-{r}-MONO-01",
            pset="S-R", manual="Tek seferde bas, 3B tara. Wattmetreyle ölç (T9 SCR tabanı).", seam=seam, layer="")
    for model, label in [("G7", "G7"), ("R2", "R2 roket"), ("X1", "ek model 1"), ("X2", "ek model 2"), ("X3", "ek model 3")]:
        for sup in ["normal", "ağaç"]:
            for k in range(1, 3):
                sid = f"T7-{model}-{'n' if sup == 'normal' else 't'}-{k:02d}"
                g = model if model in MODEL else label
                add("F5 T5–T8", "T7", g, "INS-D", "PLA", f"destek: {sup}", sid, CSTAR, "", "S-G",
                    (ins_cmd(model, sid, 100, "<C*>") if model == "G7" else
                     f"python benchmark/run_insert.py {model}.gcode -o {sid}.gcode --part-layer <dikiş> --clearance <C*>"),
                    f"Bambu Studio'da destek tipi: {sup}. " + (BOTTOM.format(n=101) if model == "G7" else "Alt parçayı dikişte durdur."),
                    seam=None if model == "G7" else "belirlenecek", layer=None if model == "G7" else "")

    # Phase 6: T9, T11, T12
    for model in ["R1", "R2"]:
        for pct in [25, 50, 75, 90]:
            sid = f"T9-{model}-h{pct}"
            add("F6 T9–T12", "T9", model, "INS-D", "PLA", f"arıza yüksekliği %{pct}", sid, CSTAR, "", "S-R",
                f"python benchmark/run_insert.py {model}.gcode -o {sid}.gcode --part-layer <%{pct} katmanı> --clearance <C*>",
                "Alt parça + insert baskılarının ikisini de wattmetre ve teraziyle ölç; süreleri kaydet.",
                seam=f"%{pct}", layer="")
    for op in ["O1", "O2", "O3"]:
        for g in ["G1", "G2"]:
            for c in [CSTAR, "0.20"]:
                for k in range(1, 3):
                    sid = f"T11-{op}-{g}-c{'S' if c == CSTAR else '020'}-{k:02d}"
                    add("F6 T9–T12", "T11", g, "INS-D", "PLA", f"operatör {op}, boşluk {c}", sid, c, "", "S-G",
                        ins_cmd(g, sid, 100, "<C*>" if c == CSTAR else c),
                        BOTTOM.format(n=101) + " Operatör yalnızca protokol kartını kullanır.")
    for rh in ["yok", "var"]:
        for k in range(1, 4):
            sid = f"T12-M2-reheat{rh[0]}-{k:02d}"
            add("F6 T9–T12", "T12", "M2", "INS-D", "PLA", f"yeniden ısıtma: {rh}", sid, CSTAR, "", "S-M",
                ins_cmd("M2", sid, 200, "<C*>") + (" <yeniden ısıtma bayrağı>" if rh == "var" else ""),
                BOTTOM.format(n=201) + " Termal kamerayla dikiş yüzeyini kaydet. Yeniden ısıtma özelliği eklenince yapılır.")
    return out


BLOCK = {"F0 Pilot": 3, "F1 T1A+T2": 8, "F2 T1B+MAN": 8, "F3 T4": 8, "F4 T3": 3, "F5 T5–T8": 6, "F6 T9–T12": 6}
WEEKS = {"F0 Pilot": "1–2", "F1 T1A+T2": "3–5", "F2 T1B+MAN": "6", "F3 T4": "6–7", "F4 T3": "8–10",
         "F5 T5–T8": "10–12", "F6 T9–T12": "12–13"}


def build():
    rng = random.Random(SEED)
    data = rows()
    ordered = []
    for phase in BLOCK:
        chunk = [r for r in data if r["phase"] == phase]
        rng.shuffle(chunk)
        for i, r in enumerate(chunk):
            r["block"] = f"{phase.split()[0]}-B{i // BLOCK[phase] + 1:02d}"
        ordered += chunk

    wb = Workbook()
    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="D0D0D0")
    border = Border(bottom=thin)

    # --- Oku beni ---------------------------------------------------------------
    ws0 = wb.active
    ws0.title = "Oku beni"
    lines = [
        ("LayerRescue insert mode — test çalışma listesi", True),
        (f"Randomizasyon tohumu (seed): {SEED}. Liste benchmark/make_run_plan.py ile yeniden üretilebilir.", False),
        ("", False),
        ("Nasıl kullanılır", True),
        ("1. 'Çalışma listesi'ndeki satırları Sıra numarasına göre bas. Sıra her fazın içinde rastgeledir; faz sırasını bozma.", False),
        ("2. Her insert numunesi iki baskıdır: önce alt parça (Elle adım sütunu), sonra Komut sütunundaki G-code.", False),
        ("   İlk alt parçada Bambu Studio önizlemesinde duraklamanın son sağlam katman (örn. 100) bittikten sonra geldiğini doğrula.", False),
        ("3. Sarı sütunlar senin dolduracağın yerler: durum, tarih, ortam, operatör ve ölçümler. Diğer sütunlara dokunma.", False),
        ("4. F1 bitince 'Parametreler' sayfasında C* (en iyi boşluk payı), F3 bitince R* (en iyi yapışma reçetesi) hücrelerini doldur.", False),
        ("   C* girilince bütün <C*> satırlarının Boşluk sütunu kendiliğinden güncellenir. Komuttaki <C*> yerine bu değeri yaz.", False),
        ("5. Her baskıda dilimleyici ayarlarını Parametreler sayfasındaki setten (S-G, S-M, S-R) değiştirmeden kullan.", False),
        ("6. G-code dosyalarını numune ID'siyle adlandır ve sakla (veri seti için).", False),
        ("", False),
        ("Fazlar ve neden bu sıra", True),
        ("F0 Pilot: ölçüm hatası ve çekme numunelerinin varyasyonu (CV) ölçülür; n gerekirse artırılır.", False),
        ("F1 T1A + T2: en iyi boşluk payı (C*) seçilir. Sonraki bütün insert baskıları C* kullanır.", False),
        ("F2 T1B + ablasyon: duvar yüksekliği ve duvarsız elle hizalama karşılaştırması.", False),
        ("F3 T4: yapışma katmanı reçetesi taranır; en iyisi R* olur.", False),
        ("F4 T3: ana dayanım deneyi C* ve R* ile; T4'ün en iyi reçetesi varsayılana karşı doğrulanır.", False),
        ("F5 T5–T8 ve F6 T9–T12: hata duyarlılığı, gerçek modeller, destekler, kaynak, operatör, termal.", False),
        ("T10 (yazılım korpusu) baskı gerektirmez; bu listede yok, F0–F1 sırasında paralel yürütülür.", False),
        ("", False),
        ("Örnek doldurulmuş satır (gerçek veri değildir)", True),
        ("Durum: Ölçüldü · Tarih: 2026-10-20 · Ortam: 23,4 °C, %48 · Operatör: O1 · Başarı: E · Olay: yok · "
         "Δx 0,06 · Δy −0,04 · θ 0,2 · e_xy otomatik hesaplanır (0,07).", False),
    ]
    for i, (t, b) in enumerate(lines, start=1):
        c = ws0.cell(row=i, column=1, value=t)
        c.font = Font(name="Arial", size=13 if i == 1 else 10, bold=b)
    ws0.column_dimensions["A"].width = 130

    # --- Parametreler -----------------------------------------------------------
    ws1 = wb.create_sheet("Parametreler")
    ws1["A1"], ws1["A1"].font = "Deneyle belirlenen değerler (doldur)", Font(name="Arial", size=12, bold=True)
    ws1["A2"], ws1["B2"], ws1["C2"] = "C*", "En iyi boşluk payı (mm), T1A sonucu", None
    ws1["A3"], ws1["B3"], ws1["C3"] = "R*", "En iyi yapışma reçetesi (T4 koşu kodu, örn. r06)", None
    for r in (2, 3):
        ws1.cell(row=r, column=3).fill = input_fill
        for col in (1, 2, 3):
            ws1.cell(row=r, column=col).font = font
    ws1["C2"].number_format = "0.00"
    ws1["D3"] = "=IFERROR(\"--temp-boost \"&INDEX(C22:C29,MATCH(C3,B22:B29,0))&\" --speed \"&INDEX(D22:D29,MATCH(C3,B22:B29,0))&IF(INDEX(E22:E29,MATCH(C3,B22:B29,0))=\"on\",\" --fan-on\",\" --fan-off\"),\"R* girilince bayraklar burada çıkar\")"
    ws1["D3"].font = font

    ws1["A5"], ws1["A5"].font = "Dilimleyici ayar setleri (Bambu Studio, P1S, 0,4 mm nozul)", Font(name="Arial", size=12, bold=True)
    sets = [
        ("Set", "Kullanım", "Katman", "Duvar", "Dolgu", "Üst/alt", "Brim", "Malzeme profili", "Not"),
        ("S-M", "M1, M2 (mekanik)", "0,20 mm", "4", "%100 çizgisel ±45°", "5 / 5", "M1: 10 mm", "Bambu PLA Basic / PETG Basic (lot kaydet)", "M1: dış duvar 60 mm/s, hızlanma 2000–3000 mm/s², en kısa katman süresi 8–10 s, Z-hop açık"),
        ("S-G", "G1–G7, T2, T11", "0,20 mm", "4", "%20 gyroid", "5 / 5", "yok", "Bambu PLA Basic", "Nervürler için 'Detect thin wall' açık"),
        ("S-R", "R1–R4, X modelleri", "0,20 mm", "profil", "profil", "profil", "profil", "Testteki malzeme", "Mevcut test 3MF ayarları korunur"),
    ]
    for i, row in enumerate(sets, start=6):
        for j, v in enumerate(row, start=1):
            c = ws1.cell(row=i, column=j, value=v)
            c.font = bold if i == 6 else font

    ws1["A11"], ws1["A11"].font = "LayerRescue insert varsayılanları (aksi yazmadıkça)", Font(name="Arial", size=12, bold=True)
    defaults = [
        ("Seçenek", "Değer", "Bayrak"),
        ("Duvar yüksekliği", "clamp(Z/2, 3, 15), en fazla Z−1", "--wall-height"),
        ("Duvar kalınlığı", "4 çizgi", "--wall-lines 4"),
        ("Duvar brimi", "5 mm", "--brim 5"),
        ("Duraklama sıcaklığı", "140 °C", "--standby-temp 140"),
        ("Yapışma katmanları", "2 katman, +10 °C, %50 hız, fan kapalı", "--adhesion-layers 2 --temp-boost 10 --speed 50 --fan-off"),
        ("Destek yeniden basımı", "açık", "(kapatmak: --no-reprint-supports)"),
        ("Z ince ayar", "0 (en fazla ±0,3 mm)", "--z-fine"),
    ]
    for i, row in enumerate(defaults, start=12):
        for j, v in enumerate(row, start=1):
            c = ws1.cell(row=i, column=j, value=v)
            c.font = bold if i == 12 else font

    ws1["A20"], ws1["A20"].font = "T4 kesirli faktoriyel 2^(4−1), E = ABC (D: yeniden ısıtma, özellik eklenince 2^(5−1)'e çıkar)", Font(name="Arial", size=12, bold=True)
    hdr = ("", "Koşu", "A: sıcaklık artışı (°C)", "B: hız (%)", "C: parça fanı", "E: bekleme (dk)")
    for j, v in enumerate(hdr, start=1):
        ws1.cell(row=21, column=j, value=v).font = bold
    for i, r in enumerate(T4_RUNS, start=22):
        for j, v in enumerate(("", r["run"], r["A"], r["B"], r["C"], r["E"]), start=1):
            ws1.cell(row=i, column=j, value=v).font = font
    for col, w in zip("ABCDEFGHI", (14, 36, 22, 24, 22, 16, 12, 38, 40)):
        ws1.column_dimensions[col].width = w

    # --- Çalışma listesi --------------------------------------------------------
    ws = wb.create_sheet("Çalışma listesi")
    cols = [
        ("Sıra", 6), ("Faz", 13), ("Hafta", 7), ("Blok (gün)", 10), ("Numune ID", 24), ("Test", 11), ("Model", 26),
        ("Grup", 11), ("Malzeme", 8), ("Değişken", 30), ("Dikiş Z (mm)", 10), ("--part-layer", 10),
        ("Boşluk (mm)", 10), ("Duvar (mm)", 9), ("Ayar seti", 10), ("Komut", 70), ("Elle adım", 60),
        ("Durum", 11), ("Tarih", 11), ("Ortam °C", 8), ("Nem %", 7), ("Operatör", 9), ("Başarı (E/H)", 9),
        ("Olay", 18), ("Δx (mm)", 8), ("Δy (mm)", 8), ("θ (°)", 7), ("e_xy (mm)", 9), ("Basamak (mm)", 10),
        ("Kuvvet (N)", 9), ("σmax (MPa)", 10), ("E (GPa)", 8), ("ε (%)", 7), ("Kırılma yeri (mm)", 11),
        ("Filament (g)", 10), ("Süre (dk)", 9), ("Enerji (Wh)", 10), ("Not", 30),
    ]
    for j, (name, w) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=name)
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 30
    first_input = 18
    for i, r in enumerate(ordered, start=2):
        clearance = r["clearance"]
        if clearance == CSTAR:
            clearance = "=IF(Parametreler!$C$2=\"\",\"C*\",Parametreler!$C$2)"
        elif clearance:
            clearance = float(clearance)
        vals = [i - 1, r["phase"], WEEKS[r["phase"]], r["block"], r["sid"], r["test"], r["model"], r["group"],
                r["material"], r["variable"], r["seam"], r["layer"], clearance, float(r["wall"]) if r["wall"] else "",
                r["pset"], r["cmd"], r["manual"], "Planlandı"]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = font
            c.border = border
        ws.cell(row=i, column=28, value=f"=IF(AND(ISNUMBER(Y{i}),ISNUMBER(Z{i})),SQRT(Y{i}^2+Z{i}^2),\"\")").font = font
        ws.cell(row=i, column=28).number_format = "0.000"
        for j in range(first_input, len(cols) + 1):
            if j != 28:
                ws.cell(row=i, column=j).fill = input_fill
                ws.cell(row=i, column=j).font = font
    last = len(ordered) + 1
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    dv = DataValidation(type="list", formula1='"Planlandı,Basıldı,Ölçüldü,Başarısız,İptal"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"E,H"', allow_blank=True)
    ws.add_data_validation(dv)
    ws.add_data_validation(dv2)
    dv.add(f"R2:R{last}")
    dv2.add(f"W2:W{last}")
    ws.conditional_formatting.add(f"A2:Q{last}", FormulaRule(formula=['$R2="Ölçüldü"'], fill=PatternFill("solid", fgColor="E2EFDA")))
    ws.conditional_formatting.add(f"A2:Q{last}", FormulaRule(formula=['OR($R2="Başarısız",$R2="İptal")'], fill=PatternFill("solid", fgColor="F8CBAD")))
    ws["P1"].comment = Comment("<C*> ve <R* bayrakları> yerine Parametreler sayfasındaki C2 ve D3 değerlerini yaz.", "plan")

    # --- Özet -------------------------------------------------------------------
    ws2 = wb.create_sheet("Özet")
    hdr = ("Faz", "Hafta", "Planlanan baskı", "Ölçüldü", "Başarısız/iptal", "İlerleme")
    for j, v in enumerate(hdr, start=1):
        c = ws2.cell(row=1, column=j, value=v)
        c.font, c.fill = head_font, head_fill
    rng_b = f"'Çalışma listesi'!$B$2:$B${last}"
    rng_r = f"'Çalışma listesi'!$R$2:$R${last}"
    for i, ph in enumerate(BLOCK, start=2):
        ws2.cell(row=i, column=1, value=ph)
        ws2.cell(row=i, column=2, value=WEEKS[ph])
        ws2.cell(row=i, column=3, value=f"=COUNTIF({rng_b},A{i})")
        ws2.cell(row=i, column=4, value=f"=COUNTIFS({rng_b},A{i},{rng_r},\"Ölçüldü\")")
        ws2.cell(row=i, column=5, value=f"=COUNTIFS({rng_b},A{i},{rng_r},\"Başarısız\")+COUNTIFS({rng_b},A{i},{rng_r},\"İptal\")")
        ws2.cell(row=i, column=6, value=f"=IF(C{i}=0,\"\",D{i}/C{i})")
        ws2.cell(row=i, column=6).number_format = "0%"
    t = len(BLOCK) + 2
    ws2.cell(row=t, column=1, value="Toplam").font = bold
    for col in "CDE":
        ws2[f"{col}{t}"] = f"=SUM({col}2:{col}{t - 1})"
    ws2[f"F{t}"] = f"=IF(C{t}=0,\"\",D{t}/C{t})"
    ws2[f"F{t}"].number_format = "0%"
    for row in ws2.iter_rows(min_row=2, max_row=t):
        for c in row:
            if c.font != bold:
                c.font = Font(name="Arial", size=10, bold=c.row == t)
    for col, w in zip("ABCDEF", (16, 9, 16, 10, 15, 10)):
        ws2.column_dimensions[col].width = w

    wb.move_sheet("Özet", offset=-2)
    wb.save(OUT)
    return ordered


if __name__ == "__main__":
    o = build()
    print(len(o), "satır ->", OUT)
